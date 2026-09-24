"""Synthetic documents for the native intake integration test; no user files."""
from pathlib import Path
import sys
import fitz
from docx import Document
from pptx import Presentation

root = Path(sys.argv[1])
text = "Teamwork validation: assign an owner, share blockers, record lessons."
document = Document()
document.add_heading("Teamwork validation", 0)
document.add_paragraph(text)
document.save(root / "sample.docx")
presentation = Presentation()
slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text = "Teamwork validation"
slide.placeholders[1].text = text
presentation.save(root / "sample.pptx")
pdf = fitz.open()
page = pdf.new_page()
page.insert_text((50, 80), text)
pdf.save(root / "sample.pdf")
pdf.close()
