"""Synthetic acceptance input, not a customer proposal or verified product data."""
from pathlib import Path
import sys
from docx import Document

root = Path(sys.argv[1])
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading("园区网络升级：先治理，再扩容", 0)
doc.add_paragraph("网站验收用虚构需求材料。下列数据均为测试假设，不代表真实客户或厂商性能。不得添加设备型号、产品参数、报价或成功案例。")
doc.add_heading("现状与目标", 1)
doc.add_paragraph("一个园区有教学区、办公区、访客区。现网存在三个问题：身份边界不清、会议时段无线拥堵、变更缺少回退记录。此次目标是先明确安全边界与变更责任，再依据实测负载扩容，不追求一次性全部替换。")
doc.add_heading("访问策略", 1)
table = doc.add_table(rows=1, cols=3)
for cell, value in zip(table.rows[0].cells, ["区域", "允许访问", "限制"]):
    cell.text = value
for row in [
    ["教学区", "教学应用和互联网", "不能直接访问办公管理网"],
    ["办公区", "授权办公系统和互联网", "按角色授权管理操作"],
    ["访客区", "互联网", "不得访问内部业务系统"],
]:
    for cell, value in zip(table.add_row().cells, row):
        cell.text = value
doc.add_heading("升级顺序与责任", 1)
doc.add_paragraph("第一步：盘点。运维负责人核对终端、链路和应用清单，形成当前基线。第二步：试点。网络负责人先在一个办公楼层实施身份隔离和无线优化，并记录优化前后结果。第三步：推广。项目负责人审核试点结果和回退预案后，分区域安排维护窗口。")
doc.add_heading("验收与回退", 1)
doc.add_paragraph("验收检查四项：访客无法访问内网；教学和办公按授权访问；对同一会议区域的优化前后体验进行实测对比；每次变更具备责任人、维护窗口和回退步骤。发现核心应用不可用时停止推广，恢复上一份配置，由网络负责人复核原因。")
doc.add_heading("待确认事项", 1)
doc.add_paragraph("当前并发终端数、链路峰值负载、既有设备可用能力和业务维护窗口均未确认。不能在这些数据缺失时直接确定设备型号、数量或预算。")
doc.save(root / "园区网络升级-验收材料.docx")
